from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from flask import send_file
import io
from datetime import datetime

from ThuVienSo.data.models.borrow_record import BorrowRecord
from ThuVienSo.data.models.book import Book


def export_report_excel():

    wb = Workbook()

    # ===================== SHEET 1: SUMMARY =====================
    ws1 = wb.active
    ws1.title = "Tổng quan"

    title_font = Font(size=14, bold=True)
    header_fill = PatternFill("solid", fgColor="DDDDDD")

    ws1["A1"] = "BÁO CÁO THƯ VIỆN"
    ws1["A1"].font = Font(size=16, bold=True)
    ws1.merge_cells("A1:B1")

    total = BorrowRecord.query.count()
    borrowing = BorrowRecord.query.filter_by(status="borrowing").count()
    returned = BorrowRecord.query.filter_by(status="returned").count()
    overdue = BorrowRecord.query.filter(
        BorrowRecord.status == "borrowing",
        BorrowRecord.due_date < datetime.utcnow()
    ).count()

    summary_data = [
        ("Tổng lượt mượn", total),
        ("Đang mượn", borrowing),
        ("Đã trả", returned),
        ("Trễ hạn", overdue),
    ]

    row = 3
    for label, value in summary_data:
        ws1[f"A{row}"] = label
        ws1[f"B{row}"] = value
        ws1[f"A{row}"].font = title_font
        row += 1

    # ===================== SHEET 2: BORROW DETAIL =====================
    ws2 = wb.create_sheet("Chi tiết mượn")

    ws2.append(["ID", "User", "Ngày mượn", "Hạn trả", "Trạng thái"])

    records = BorrowRecord.query.all()

    for r in records:
        ws2.append([
            r.id,
            r.user.full_name if r.user else "",
            str(r.borrow_date),
            str(r.due_date),
            r.status
        ])

    # ===================== SHEET 3: BOOK STOCK =====================
    ws3 = wb.create_sheet("Kho sách")

    ws3.append(["ID", "Tên sách", "Tổng", "Còn lại"])

    books = Book.query.all()

    for b in books:
        total_qty = sum(c.total_quantity for c in b.copies)
        available_qty = sum(c.available_quantity for c in b.copies)

        ws3.append([
            b.id,
            b.title,
            total_qty,
            available_qty
        ])

    # ===================== OUTPUT =====================
    file = io.BytesIO()
    wb.save(file)
    file.seek(0)

    return send_file(
        file,
        as_attachment=True,
        download_name="library_report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )